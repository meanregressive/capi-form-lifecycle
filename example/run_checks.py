#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
run_checks.py — regression test for the skill's scripts, using the starter example as fixture.

    python example/run_checks.py

1. regenerates example/ from _generate_example.py (so the fixture is always current);
2. validate_xlsform.py on the Round-1 form must PASS with 0 FAIL;
3. extract_for_diff.py must read the EN paper with the tracked change ACCEPTED and the two-comment
   thread linked (reply -> parent), and the PT paper with the planted "por mês";
4. paper_from_xlsform.py --batch must write one docx per language with the same number of tables;
5. build_diff_docx.py must assemble expected/findings_01.md into a report with 6 findings rows.
Exit code 1 on any failure.
"""
import os, re, subprocess, sys, shutil, tempfile
from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE); S = os.path.join(ROOT, "scripts")
PY = sys.executable; fails = []

def run(*args, ok=(0,)):
    r = subprocess.run([PY, *args], capture_output=True, text=True, encoding="utf-8", errors="replace")
    if r.returncode not in ok: fails.append(f"{os.path.basename(args[0])} exit {r.returncode}\n{r.stdout[-800:]}\n{r.stderr[-800:]}")
    return r

def check(cond, msg):
    print(("[PASS] " if cond else "[FAIL] ") + msg)
    if not cond: fails.append(msg)

run(os.path.join(HERE, "_generate_example.py"))
form = os.path.join(HERE, "03_input_prior_capi", "1_school_head_r1_v1.xlsx")
r = run(os.path.join(S, "validate_xlsform.py"), form, "--attachments", os.path.join(HERE, "07_capi_attachments", "1_school_head"))
check("PASSED: 0 FAIL" in r.stdout, "validate_xlsform: Round-1 form passes")

tmp = tempfile.mkdtemp(prefix="capi_example_")
run(os.path.join(S, "extract_for_diff.py"), "--out", tmp, "--name", "en_01", os.path.join(HERE, "01_input_paper_EN", "20260105", "1_school_head_EN.docx"))
run(os.path.join(S, "extract_for_diff.py"), "--out", tmp, "--name", "pt_01", os.path.join(HERE, "02_input_paper_PT", "20260105", "1_school_head_PT.docx"))
run(os.path.join(S, "extract_for_diff.py"), "--out", tmp, "--name", "capi_01", form)
en = open(os.path.join(tmp, "en_01.txt"), encoding="utf-8").read(); pt = open(os.path.join(tmp, "pt_01.txt"), encoding="utf-8").read()
capi = open(os.path.join(tmp, "capi_01.txt"), encoding="utf-8").read()
check("enrolled this school year" in en and "are enrolled? |" not in en, "extract: tracked change read as accepted")
check("(reply to C0)" in en, "extract: comment reply linked to its parent")
check("por mês" in pt, "extract: PT misalignment present")
check("[select_one yesno] consent" in capi and "form_id: school_head_r1_v1" in capi, "extract: XLSForm dump has survey rows and settings")

run(os.path.join(S, "paper_from_xlsform.py"), "--batch", os.path.join(HERE, "05_python_scripts", "paper_configs", "batch.json"))
pe = os.path.join(HERE, "06_output_paper_FROMCAPI", "EN", "1_school_head_EN_FROMCAPI_r1.docx")
pp = os.path.join(HERE, "06_output_paper_FROMCAPI", "PT", "1_school_head_PT_FROMCAPI_r1.docx")
if os.path.exists(pe) and os.path.exists(pp):
    de, dp = Document(pe), Document(pp)
    check(len(de.tables) == len(dp.tables) and len(de.tables) >= 4, f"paper: EN/PT same table count ({len(de.tables)})")
    txt = "\n".join(c.text for t in de.tables for r in t.rows for c in r.cells)
    check("[SHOW IF: ${C02}=1]" in txt and "Prefilled" in txt, "paper: relevance markers and Prefilled cells present")
    check(not re.search(r"\$\{(?!C0|E0|B0|D0|A0|consent)[A-Za-z_]+\}", "\n".join(p.text for p in de.paragraphs)), "paper: no leaked variables in body text")
    for sub in ("EN", "PT"):
        os.makedirs(os.path.join(HERE, "expected", "paper", sub), exist_ok=True)
    shutil.copyfile(pe, os.path.join(HERE, "expected", "paper", "EN", os.path.basename(pe)))
    shutil.copyfile(pp, os.path.join(HERE, "expected", "paper", "PT", os.path.basename(pp)))
else:
    check(False, "paper: batch wrote both language files")

# formatting carry-through: synthetic patch via xlsform_io (label change + insert + delete), then compare
sys.path.insert(0, S)
try:
    from xlsform_io import XLSFormBook, summary
    before = summary(form)["survey"]
    book = XLSFormBook(form); hdr, rows = book.header("survey"), book.rows("survey")
    ci = {h: i for i, h in enumerate(hdr)}
    names = [r[ci["name"]] for r in rows]
    rows[names.index("B03")][ci["hint:ENG"]] = "0 if under one year."                         # LABEL/HINT change
    new = [None] * len(hdr); new[ci["type"]] = "select_one yesno"; new[ci["name"]] = "B06"
    new[ci["label:ENG"]] = "B06. Does your phone have internet access?"; new[ci["label"]] = "B06. O seu telefone tem acesso à internet?"
    new[ci["required"]] = "yes"; rows.insert(names.index("B05_o") + 1, new)                    # INSERT
    del rows[[r[ci["name"]] for r in rows].index("E02")]                                     # DELETE
    patched = os.path.join(tmp, "patched.xlsx"); rep = book.save(patched); after = summary(patched)["survey"]
    check(after["freeze"] == before["freeze"] == "A2", "xlsform_io: frozen header kept")
    check(after["widths"] == before["widths"] and after["hidden_cols"] == before["hidden_cols"], "xlsform_io: column widths and hidden column kept")
    check(after["fills"] == before["fills"] and after["colored_font"] == before["colored_font"], f"xlsform_io: fills and colored fonts kept ({after['fills']} fills)")
    check(after["cf_rules"] == before["cf_rules"] and after["dv"] == before["dv"] and after["tab"] == before["tab"], "xlsform_io: conditional formatting, validation and tab color kept")
    wb2 = __import__("openpyxl").load_workbook(patched); ws2 = wb2["survey"]; last = ws2.max_row
    ours = [str(cf.sqref) for cf in ws2.conditional_formatting if any('<>""' in f for r in cf.rules for f in (r.formula or []))]
    check(ours and ours[0].endswith(str(last)), f"xlsform_io: example rule re-anchored to new last row ({ours[0] if ours else 'missing'})")
    tpl_rules = sum(len(cf.rules) for cf in ws2.conditional_formatting)
    hp = {c.value: i for i, c in enumerate(ws2[1], start=1)}; L = __import__("openpyxl").utils.get_column_letter
    text_rule = [str(cf.sqref) for cf in ws2.conditional_formatting if any(r.formula == ['$A1="text"'] for r in cf.rules)]
    widened = any(t.startswith(f"B1:{L(hp['hint'])}") for t in text_rule)   # name..hint now spans the inserted ENG columns
    check(tpl_rules == 32 and widened, f"xlsform_io: template's 31 type-color rules kept as 31 (+1 ours = {tpl_rules}) and widened to the bilingual columns ({text_rule})")
    check(all(s in wb2.sheetnames for s in ("help-survey", "help-choices", "help-settings")), "xlsform_io: template help sheets ride along untouched")
    r_b06 = [r for r in range(2, last + 1) if ws2.cell(r, ci["name"] + 1).value == "B06"][0]
    check(ws2.cell(r_b06, ci["label"] + 1).fill.fgColor.rgb == ws2.cell(r_b06 - 1, ci["label"] + 1).fill.fgColor.rgb and rep["rows_inherited"] == 1, "xlsform_io: inserted row inherited the style above it")
    vals = [ws2.cell(r, ci["name"] + 1).value for r in range(2, last + 1)]
    check("E02" not in vals and "B06" in vals and ws2.max_column == len(hdr), "xlsform_io: values rebuilt correctly, no stale columns")
    r2 = run(os.path.join(S, "validate_xlsform.py"), patched, "--base", form, "--planned", "B03,B06,E02")
    check("intended-diff" in r2.stdout and "0 changed fields outside the plan" in r2.stdout, "validate: intended-diff passes on the xlsform_io patch")
except Exception as e:
    check(False, f"xlsform_io: {type(e).__name__}: {e}")

pdf = os.path.join(HERE, "01_input_paper_EN", "20260105", "1_school_head_EN.pdf")
if os.path.exists(pdf):
    run(os.path.join(S, "extract_for_diff.py"), "--out", tmp, "--name", "en_01_pdf", pdf)
    dumped = os.path.join(tmp, "en_01_pdf.txt")
    if os.path.exists(dumped):
        d = open(dumped, encoding="utf-8").read()
        check(d.startswith("# PDF-DERIVED"), "pdf fallback: dump is marked PDF-DERIVED")
        check("enrolled this school year" in d and "| C04 |" in d, "pdf fallback: question rows survive PDF -> docx -> dump")
        check(os.path.exists(os.path.join(tmp, "en_01_pdf_FROMPDF.docx.conversion_report.txt")), "pdf fallback: conversion report written")
    else:
        check(False, "pdf fallback: extractor accepted the .pdf")
else:
    print("[SKIP] pdf fallback: no example PDF fixture present")

rep = os.path.join(tmp, "report.docx")
run(os.path.join(S, "build_diff_docx.py"), "--dumps", os.path.join(HERE, "expected"), "--out", rep, "--project", "DEMO", "--round", "Round 2")
if os.path.exists(rep):
    d = Document(rep); rows = len(d.tables[1].rows) - 1 if len(d.tables) > 1 else 0
    check(rows == 6, f"build_diff_docx: findings table has 6 rows (got {rows})")
else:
    check(False, "build_diff_docx: report written")

shutil.rmtree(tmp, ignore_errors=True)
print("\n" + ("FAILED" if fails else "ALL CHECKS PASSED") + f": {len(fails)} failure(s)")
for f in fails: print(" -", f)
sys.exit(1 if fails else 0)
