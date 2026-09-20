#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_diff_docx.py — assemble per-instrument findings_NN.md files (DIFF mode, Phase D2) into ONE
landscape Word report, color-coded by severity, with an executive summary table.

Usage
-----
  python build_diff_docx.py --dumps <dumps_dir> --out <report.docx>
                            --project "<Project name>" --round "<Round label>"
                            [--instruments 01,02,09] [--lang-tag EN] [--lang2-tag PT]
                            [--method "<one paragraph replacing the default method text>"]

  --dumps        folder holding findings_NN.md (and optionally findings_NN_<YYYYMMDD>.md; the
                 newest file per instrument number is used)
  --instruments  comma-separated register numbers to include, in order; default = every
                 findings file found, sorted by number
  --lang-tag     tag of the design language as written in the findings tables (default EN)
  --lang2-tag    tag of the field language (default PT); used only in legend text

Expected findings_NN.md structure (fixed by modes/diff.md Phase D2):
  # Instrument NN — <name>
  SOURCES: CAPI=<...> ; EN=<...> ; PT=<...>
  (optional prose notes)
  ## DIFFERENCES
  | Q code | Section / topic | Diff type | Paper says | CAPI has | Severity | Recommendation |
  ## SUMMARY
  - bullets

Nothing project-specific is hard-coded: names come from the findings titles, sources from the
SOURCES lines, the title block from --project / --round. Requires python-docx.
"""
import argparse, glob, os, re, sys
from datetime import date

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:  # pragma: no cover
    raise SystemExit("python-docx is required: pip install python-docx")

SEV_FILL = {"High": "F4CCCC", "Med": "FCE5CD", "Low": "EFEFEF", "Doc": "D9E2F3", "": "FFFFFF"}
HEADER_FILL = "1F3864"
COL_W = [0.6, 1.2, 0.95, 2.6, 2.2, 0.6, 1.85]  # inches, sums to 10.0 (landscape letter, 0.5" margins)


# ----------------------------------------------------------------------------- docx helpers
def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)

def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true")
    trPr.append(th)

def set_cell_text(cell, text, size=8.5, bold=False, color=None, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None: p.alignment = align
    p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
    run = p.add_run(text); run.font.size = Pt(size); run.font.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)
    return cell

def h(doc, text, size=14, color="1F3864", space_before=10, space_after=4, bold=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before); p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text); r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = RGBColor.from_string(color)
    return p

def body(doc, text, size=9.5, italic=False, bold=False, space_after=3):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text); r.font.size = Pt(size); r.italic = italic; r.font.bold = bold
    return p

def bullet(doc, text, size=9.0):
    p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(1)
    r = p.add_run(text); r.font.size = Pt(size)
    return p

def sev_class(text):
    t = text or ""
    if "High" in t: return "High"
    if "Med" in t: return "Med"
    if "Low" in t: return "Low"
    if "doc" in t.lower(): return "Doc"
    return ""


# ----------------------------------------------------------------------------- findings parser
def parse_findings(path):
    with open(path, encoding="utf-8") as f:
        lines = [l.rstrip("\n") for l in f]
    title, sources, notes, rows, summary = "", "", [], [], []
    section = "head"
    for ln in lines:
        s = ln.strip()
        if s.startswith("# Instrument"):
            title = s.lstrip("# ").strip(); continue
        if s.upper().startswith("SOURCES:"):
            sources = s[len("SOURCES:"):].strip(); continue
        if s.startswith("## DIFFERENCES"): section = "table"; continue
        if s.startswith("## SUMMARY"): section = "summary"; continue
        if section == "head":
            if s and not s.startswith("#"): notes.append(s)
            continue
        if section == "table":
            if s.startswith("|"):
                cells = [c.strip() for c in s.strip("|").split("|")]
                if all(c and set(c) <= set("-: ") for c in cells): continue          # separator row
                if cells and cells[0].lower().replace("-", " ") in ("q code", "qcode"): continue  # header row
                if len(cells) < 7: cells += [""] * (7 - len(cells))
                elif len(cells) > 7: cells = cells[:6] + [" / ".join(cells[6:])]
                rows.append(cells)
            continue
        if section == "summary":
            if s.startswith(("-", "*")): summary.append(s.lstrip("-* ").strip())
            elif s and summary: summary[-1] += " " + s
    name = title.split("—", 1)[1].strip() if "—" in title else title
    return {"title": title, "name": name, "sources": sources, "note": " ".join(notes).strip(),
            "rows": rows, "summary": summary, "path": path}

def find_findings(dumps, instruments):
    """Return [(NN, path)] — newest findings file per instrument number."""
    found = {}
    for p in glob.glob(os.path.join(dumps, "findings_*.md")):
        m = re.match(r"findings_(\d+)(?:_(\d{8}))?\.md$", os.path.basename(p))
        if not m: continue
        num, stamp = m.group(1), m.group(2) or ""
        if num not in found or stamp > found[num][0]: found[num] = (stamp, p)
    order = instruments or sorted(found, key=lambda x: int(x))
    out = []
    for num in order:
        num = num.zfill(2) if len(num) < 2 else num
        if num in found: out.append((num, found[num][1]))
        else: print(f"WARN  no findings file for instrument {num} in {dumps}")
    return out


# ----------------------------------------------------------------------------- tables
def add_findings_table(doc, rows, lang_tag, lang2_tag):
    head = ["Q code", "Section / topic", "Diff type", f"Paper says ({lang_tag}; note {lang2_tag} if differs)",
            "CAPI has", "Severity", "Recommendation"]
    table = doc.add_table(rows=1, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed"); table._tbl.tblPr.append(layout)
    hdr = table.rows[0]; set_repeat_header(hdr)
    for i, c in enumerate(hdr.cells):
        set_cell_text(c, head[i], size=8.5, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_bg(c, HEADER_FILL)
    for r in rows:
        cells = table.add_row().cells
        for i, val in enumerate(r):
            align = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 5) else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[i], val, size=8.0, bold=(i == 5), align=align)
        set_cell_bg(cells[5], SEV_FILL.get(sev_class(r[5]), "FFFFFF"))
    for row in table.rows:
        for i, c in enumerate(row.cells): c.width = Inches(COL_W[i])
    return table

def add_exec_summary(doc, parsed):
    ex = doc.add_table(rows=1, cols=6)
    ex.style = "Table Grid"; ex.alignment = WD_TABLE_ALIGNMENT.CENTER
    exhdr = ["#", "Instrument", "Sources compared", "High", "Med", "Low"]
    exw = [0.4, 2.3, 6.1, 0.4, 0.4, 0.4]
    for i, c in enumerate(ex.rows[0].cells):
        set_cell_text(c, exhdr[i], size=9, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_bg(c, HEADER_FILL)
    tot = {"High": 0, "Med": 0, "Low": 0}
    for num, d in parsed:
        counts = {"High": 0, "Med": 0, "Low": 0}
        for r in d["rows"]:
            sc = sev_class(r[5])
            if sc in counts: counts[sc] += 1
        for k in tot: tot[k] += counts[k]
        cells = ex.add_row().cells
        vals = [num.lstrip("0") or "0", d["name"], d["sources"], str(counts["High"]), str(counts["Med"]), str(counts["Low"])]
        for i, v in enumerate(vals):
            align = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 3, 4, 5) else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[i], v, size=8.5, align=align)
        if counts["High"]: set_cell_bg(cells[3], SEV_FILL["High"])
        if counts["Med"]: set_cell_bg(cells[4], SEV_FILL["Med"])
    cells = ex.add_row().cells
    for i, v in enumerate(["", "TOTAL", "", str(tot["High"]), str(tot["Med"]), str(tot["Low"])]):
        align = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 3, 4, 5) else WD_ALIGN_PARAGRAPH.LEFT
        set_cell_text(cells[i], v, size=8.5, bold=True, align=align); set_cell_bg(cells[i], "E2EFDA")
    for i, col in enumerate(ex.columns):
        for cell in col.cells: cell.width = Inches(exw[i])
    return tot


# ----------------------------------------------------------------------------- main
DEFAULT_METHOD = (
    "Each CAPI instrument (SurveyCTO XLSForm) was compared with the paper questionnaire in every language "
    "provided. Tracked changes in the paper documents were treated as accepted (final). The comparison follows "
    "the project's authority order and do-not-flag list (skill reference diff_rubric.md, project brief §4 and §6): "
    "only substantive, meaning-changing differences are listed; coding conventions, programming rows and "
    "formatting are not flagged.")

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--dumps", required=True); ap.add_argument("--out", required=True)
    ap.add_argument("--project", required=True); ap.add_argument("--round", required=True)
    ap.add_argument("--instruments", help="comma-separated register numbers in report order")
    ap.add_argument("--lang-tag", default="EN"); ap.add_argument("--lang2-tag", default="PT")
    ap.add_argument("--method", help="replace the default purpose-and-method paragraph")
    a = ap.parse_args()

    wanted = [x.strip() for x in a.instruments.split(",") if x.strip()] if a.instruments else None
    files = find_findings(a.dumps, wanted)
    if not files: raise SystemExit(f"no findings_NN.md files in {a.dumps}")
    parsed = [(num, parse_findings(p)) for num, p in files]

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"; doc.styles["Normal"].font.size = Pt(9.5)
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width, sec.page_height = Inches(11), Inches(8.5)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(0.5)

    h(doc, f"{a.project} — {a.round}", size=18, space_before=0, space_after=2)
    h(doc, "CAPI form vs paper questionnaire — differences review", size=14, color="2E74B5", space_after=6)
    body(doc, f"Prepared {date.today().isoformat()}  ·  capi-form-lifecycle DIFF mode  ·  {len(parsed)} instrument(s)",
         size=10, italic=True, space_after=8)

    h(doc, "Purpose and method", size=12)
    body(doc, a.method or DEFAULT_METHOD, space_after=6)

    h(doc, "Legend", size=12)
    body(doc, "Diff types: MISSING-IN-CAPI (paper has it, form does not) · EXTRA-IN-CAPI (content question in the form, "
              "not in the paper) · RESPONSE-OPTIONS (answer set differs) · SKIP-LOGIC · WORDING (meaning-changing) · "
              f"LANG-MISALIGN ({a.lang_tag} and {a.lang2_tag} disagree on content) · RENUMBERED/ORDER · TYPE · "
              "MISSING-DOC (a side was not provided).", size=9, space_after=2)
    pl = doc.add_paragraph(); pl.paragraph_format.space_after = Pt(2)
    for lab, col in [("High = changes or loses data", SEV_FILL["High"]), ("Med = meaning or order, low data risk", SEV_FILL["Med"]),
                     ("Low = minor, note only", SEV_FILL["Low"]), ("Doc not provided", SEV_FILL["Doc"])]:
        r = pl.add_run("  " + lab + "  "); r.font.size = Pt(9); r.font.bold = True
        rpr = r._element.get_or_add_rPr()
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), col); rpr.append(shd)
        pl.add_run("   ")

    h(doc, "Executive summary — flagged differences by instrument", size=13)
    tot = add_exec_summary(doc, parsed)
    body(doc, "Counts are flagged table rows per severity; MISSING-DOC and no-action rows are informational. "
              "See each instrument's table and summary for detail.", size=8.5, italic=True, space_after=4)

    for num, d in parsed:
        doc.add_page_break()
        h(doc, f"Instrument {num.lstrip('0') or '0'} — {d['name']}", size=15, space_before=0, space_after=2)
        if d["sources"]: body(doc, f"Sources — {d['sources']}", size=8.5, italic=True, space_after=2)
        if d["note"]: body(doc, d["note"], size=8.5, italic=True, space_after=4)
        add_findings_table(doc, d["rows"], a.lang_tag, a.lang2_tag)
        h(doc, "Summary", size=11, space_before=6, space_after=2)
        for b in d["summary"]: bullet(doc, b, size=9.0)

    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    doc.save(a.out)
    print(f"Saved: {a.out}")
    print(f"Instruments: {', '.join(n for n, _ in parsed)}  |  flagged rows High/Med/Low: {tot}")

if __name__ == "__main__":
    main()
