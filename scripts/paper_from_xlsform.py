#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
paper_from_xlsform.py — regenerate a paper questionnaire (.docx, one file per language) from a
SurveyCTO / ODK XLSForm. The engine behind PAPER mode (modes/paper.md).

Usage
-----
  One form:
    python paper_from_xlsform.py <form.xlsx> --out <dir> --base <stem>
                                 [--langs "EN=label:ENG,PT=label"] [--config <instrument.json>]
                                 [--suffix FROMCAPI_v1] [--no-autofit] [--pdf]

  Many forms (PAPER mode "all"):
    python paper_from_xlsform.py --batch <batch.json>          (batch.json may set "pdf": true)

  --pdf      also export each docx to PDF next to it, through Microsoft Word (Windows with Word and
             pywin32). Without Word the docx is still written and the PDF step says why it skipped;
             LibreOffice users can run `soffice --headless --convert-to pdf <file.docx>` instead.

  --langs    comma-separated <TAG>=<label column> pairs, in output order. The hint and
             choice-label columns are derived (label:ENG -> hint:ENG ; label -> hint).
             Default "EN=label:ENG,PT=label". Output goes to <out>/<TAG>/<base>_<TAG>_<suffix>.docx.
  --config   per-instrument JSON (all keys optional):
               {"titles":       {"EN": "...", "PT": "..."},          banner; default = form_title
                "var_labels":   {"member_name": {"EN": "this member", "PT": "este membro"}},
                "note_markers": {"consent_b": {"EN": "[Version B]", "PT": "[Versão B]"}},
                "matrices":     {"<repeat name>": {"orient": "instances_cols"|"instances_rows",
                                                    "blank_cols": 12, "header": {"EN": "Member", "PT": "Membro"},
                                                    "label_calc": {"EN": "<calc field>", "PT": "<calc field>"},
                                                    "count": 7, "corner": ""}},
                "strings":      {"FR": {...}}   UI strings for a language this script lacks (see STRINGS)}
  --batch    JSON: {"out": "<dir>", "langs": "EN=label:ENG,PT=label", "suffix": "FROMCAPI_v1",
                    "forms": [{"form": "<path.xlsx>", "base": "<stem>", "config": "<path.json>"}, ...]}

What it renders
---------------
  - 3-column table  code | question | response ; landscape; visible grid.
  - response cell:  select_* -> "value – label" lines ; calculate with pulldata -> "Prefilled" ;
                    date/datetime -> "Automatic" ; image -> "Photo" ; enumerator -> "List" ;
                    integer/decimal/text -> hint, else "Enter response".
  - groups are read recursively; a labeled group becomes a header (SECTION… 13pt, else 11pt);
    label-less wrapper groups are transparent.
  - a repeat with >= 2 displayable questions -> transposed grid, or a matrix if configured.
  - relevance / constraint are NOT turned into prose: shown as [SHOW IF: …] / [CHECK: …] markers.
  - plumbing rows (start, end, deviceid, calculate_here, internal calcs) are dropped.
  - code column: paper code from the label prefix (bold) > field name that looks like a code
    (gray italic) > "[CAPI CODE]" flag.
  - long numeric picklists collapse to a range line; roster-driven choice lists collapse to one
    line; "timed-grid" plugin fields collapse to an item count; "Other, specify" text fields fold
    into the preceding question inside matrices.

Nothing project-specific is hard-coded. Requires openpyxl and python-docx.
"""
import argparse, json, os, re, sys
import openpyxl
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.stdout = __import__("io").TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

CODE_COL_W = Inches(0.9)
LANDSCAPE_USABLE_IN = 9.5
CODE_FLAG = "[CAPI CODE]"

# UI strings the paper needs in each language. Add a language via config "strings".
STRINGS = {
    "EN": dict(item="item", enter="Enter response", prefilled="Prefilled", automatic="Automatic",
               photo="Photo", listsel="List", showif="SHOW IF", check="CHECK",
               roster_choice="select from roster", number_range="enter a number in this range",
               timed_grid="Timed grid — administered from physical card", items="items"),
    "PT": dict(item="item", enter="Introduza a resposta", prefilled="Pré-preenchido",
               automatic="Gerado automaticamente", photo="Foto", listsel="Lista",
               showif="MOSTRAR SE", check="VERIFICAR", roster_choice="selecionar da lista",
               number_range="introduza um número neste intervalo",
               timed_grid="Grelha cronometrada — administrada a partir do cartão físico", items="itens"),
    "ES": dict(item="ítem", enter="Introduzca la respuesta", prefilled="Precargado",
               automatic="Automático", photo="Foto", listsel="Lista", showif="MOSTRAR SI",
               check="VERIFICAR", roster_choice="seleccionar de la lista",
               number_range="introduzca un número en este rango",
               timed_grid="Cuadrícula cronometrada — administrada desde la tarjeta física", items="ítems"),
    "FR": dict(item="élément", enter="Saisir la réponse", prefilled="Prérempli",
               automatic="Automatique", photo="Photo", listsel="Liste", showif="AFFICHER SI",
               check="VÉRIFIER", roster_choice="sélectionner dans la liste",
               number_range="saisir un nombre dans cet intervalle",
               timed_grid="Grille chronométrée — administrée à partir de la carte physique", items="éléments"),
}

DROP_TYPES = {"start", "end", "deviceid", "subscriberid", "simserial", "phonenumber", "username",
              "calculate_here", "audio", "audio audit", "text audit", "speed violations count",
              "speed violations list", "speed violations audit", "caseid"}
NAME_AS_CODE = re.compile(r"^[A-Z]{1,2}\d{1,3}[a-z]?(?:[._]\d{1,2})?[a-z]?\d?$")
SECTION_RE = re.compile(r"\s*(SEC[ÇC][ÃA]O|SECTION|SECCI[ÓO]N|ABSCHNITT)\b", re.I)
OTHER_NAME_RE = re.compile(r"_(o|oth|other|esp|autre|otro)$", re.I)
OTHER_LBL_RE = re.compile(r"\b(other|specify|outr[oa]|especifi|autre|pr[ée]cis|otr[oa])", re.I)
PLACEHOLDER_RE = re.compile(r"^\$\{[A-Za-z0-9_]+\}$")
INT_RE = re.compile(r"^-?\d+$")


# ----------------------------------------------------------------------------- text helpers
def clean_html(s):
    if s is None: return ""
    s = str(s)
    s = re.sub(r"<p>\s*&nbsp;\s*</p>", "\n", s, flags=re.I)
    s = re.sub(r"<br\s*/?>", "\n", s, flags=re.I)
    s = re.sub(r"</p>", "\n", s, flags=re.I)
    s = re.sub(r"<[^>]+>", "", s)
    s = s.replace("&nbsp;", " ").replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r"\n{3,}", "\n\n", s)
    return s.strip()

def subst_vars(s, lang):
    if not s: return s
    s = re.sub(r"\$\{([a-z0-9_]*index|[a-z0-9_]*_idx|pos)\}", "N", s)
    for var, repl in lang["_subs"].items():
        s = re.sub(r"\$\{" + re.escape(var) + r"\}", repl, s)
    return re.sub(r"\$\{([A-Za-z0-9_]+)\}", r"[\1]", s)

def proc(s, lang):
    return subst_vars(clean_html(s), lang)

def split_code(label):
    """Pull a leading paper code ('A07.' / 'B3a.' / 'B3.1' / 'C12c2.') off the label."""
    if not label: return "", ""
    m = (re.match(r"\s*([A-Z]{1,2}\d{1,3}(?:\.\d{1,2})?[a-z]?\d?)\.\s*(.*)", label, flags=re.S)
         or re.match(r"\s*([A-Z]{1,2}\d{1,3}(?:\.\d{1,2})?[a-z]?\d?)\s+(.*)", label, flags=re.S))
    return (m.group(1), m.group(2).strip()) if m else ("", label.strip())

def code_fallback(name):
    n = (name or "").strip()
    return (n, True) if NAME_AS_CODE.match(n) else (CODE_FLAG, False)


# ----------------------------------------------------------------------------- workbook
def load(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    def sheet(name):
        rows = list(wb[name].iter_rows(values_only=True))
        idx = {str(h).strip(): i for i, h in enumerate(rows[0]) if h is not None}
        return idx, rows[1:]
    s_idx, s_rows = sheet("survey")
    c_idx, c_rows = sheet("choices") if "choices" in wb.sheetnames else ({}, [])
    title = None
    if "settings" in wb.sheetnames:
        set_idx, set_rows = sheet("settings")
        if "form_title" in set_idx and set_rows: title = set_rows[0][set_idx["form_title"]]
    wb.close()
    return s_idx, s_rows, c_idx, c_rows, title

def g(row, idx, col):
    i = idx.get(col)
    return row[i] if i is not None and i < len(row) else None

def build_choices(c_idx, c_rows):
    d = {}
    lcol = "list_name" if "list_name" in c_idx else "list name"
    for r in c_rows:
        ln = g(r, c_idx, lcol)
        if ln: d.setdefault(str(ln).strip(), []).append(r)
    return d

def any_label(row, s_idx, label_cols):
    return any(g(row, s_idx, c) for c in label_cols)


# ----------------------------------------------------------------------------- response cell
def render_options(list_name, choices, c_idx, lang):
    opts = choices.get(list_name, [])
    vcol = "value" if "value" in c_idx else "name"
    labs = [clean_html(g(r, c_idx, lang["clabel"])) for r in opts]
    vals = [g(r, c_idx, vcol) for r in opts]
    ph = sum(1 for l in labs if PLACEHOLDER_RE.match(l or ""))
    if opts and ph >= max(2, 0.8 * len(opts)):                       # dynamic roster list
        sv = [str(v) for v in vals]
        return f"{sv[0]}–{sv[-1]}: {lang['roster_choice']}"
    numeric = [(v, l) for v, l in zip(vals, labs) if INT_RE.match((l or "").strip())]
    if opts and len(numeric) >= max(5, 0.7 * len(opts)):             # numeric picklist
        nl = [int(l) for _, l in numeric]; lo, hi = min(nl), max(nl)
        items = []
        for v, l in zip(vals, labs):
            if INT_RE.match((l or "").strip()): continue
            try: sv = int(v)
            except (TypeError, ValueError): sv = 0
            items.append((sv, f"{v} – {l}"))
        items.append((lo, f"{lo}–{hi}: {lang['number_range']}"))
        items.sort(key=lambda t: (t[0] < 0, t[0]))
        return "\n".join(t for _, t in items)
    return "\n".join(f"{v} – {l}" for v, l in zip(vals, labs))

def response_cell(row, s_idx, choices, c_idx, lang):
    t = (g(row, s_idx, "type") or "").strip()
    calc = g(row, s_idx, "calculation") or ""
    hint = proc(g(row, s_idx, lang["hint"]), lang)
    appear = g(row, s_idx, "appearance") or ""
    if t.startswith(("select_one", "select_multiple")) and "timed-grid" in appear:
        ln = t.split(None, 1)[1].strip()
        cell = f"[{lang['timed_grid']}; {len(choices.get(ln, []))} {lang['items']}]"
    elif t.startswith(("select_one", "select_multiple")):
        parts = t.split(None, 1)
        cell = render_options(parts[1].strip(), choices, c_idx, lang) if len(parts) > 1 else ""
    elif t == "calculate":
        cell = lang["prefilled"] if "pulldata" in calc else None
    elif t in ("datetime", "date"): cell = lang["automatic"]
    elif t == "image": cell = lang["photo"]
    elif t == "enumerator": cell = lang["listsel"]
    elif t in ("integer", "decimal", "text"): cell = hint or lang["enter"]
    else: cell = hint or ""
    extra = []
    rel, con = g(row, s_idx, "relevance") or g(row, s_idx, "relevant"), g(row, s_idx, "constraint")
    if rel: extra.append(f"[{lang['showif']}: {rel}]")
    if con: extra.append(f"[{lang['check']}: {con}]")
    return "\n".join(p for p in [cell] + extra if p)

def displayable(row, s_idx, label_cols):
    t = (g(row, s_idx, "type") or "").strip()
    if t in DROP_TYPES: return False
    if t == "calculate":
        return "pulldata" in (g(row, s_idx, "calculation") or "") and any_label(row, s_idx, label_cols)
    if t in ("begin group", "end group", "begin repeat", "end repeat", "note", "begin_group", "end_group",
             "begin_repeat", "end_repeat"): return False
    return any_label(row, s_idx, label_cols)


# ----------------------------------------------------------------------------- docx helpers
def _tight(p, after=2):
    pf = p.paragraph_format; pf.space_before = Pt(0); pf.space_after = Pt(after); pf.line_spacing = 1.0

def add_para(doc, text, size=10, bold=False, center=False):
    for chunk in [c for c in (text or "").split("\n") if c.strip()]:
        p = doc.add_paragraph()
        if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(chunk); r.bold = bold; r.font.size = Pt(size)
        _tight(p)

def fill_cell(cell, text, size=10, bold=False, italic=False, gray=False):
    cell.text = ""
    lines = [l for l in (text or "").split("\n") if l.strip()] or [""]
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        r = p.add_run(line); r.font.size = Pt(size); r.bold = bold
        if italic: r.italic = True
        if gray: r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        _tight(p, after=0)

def fill_code_cell(cell, code, name):
    if code:
        fill_cell(cell, code, bold=True); return code
    txt, is_name = code_fallback(name)
    fill_cell(cell, txt, bold=not is_name, italic=is_name, gray=is_name)
    return txt

def set_col_width(table, col, width):
    table.autofit = False; table.allow_autofit = False
    for row in table.rows: row.cells[col].width = width


# ----------------------------------------------------------------------------- tree
class Node:
    __slots__ = ("kind", "row", "name", "label", "children")
    def __init__(self, kind, row=None, name=None, label=None):
        self.kind, self.row, self.name, self.label, self.children = kind, row, name, label, []

def parse_tree(s_rows, s_idx, lang, label_cols):
    root = Node("root"); stack = [root]
    for r in s_rows:
        if all(v is None for v in r): continue
        t = (g(r, s_idx, "type") or "").strip().replace("_", " ")
        if t in ("begin group", "begin repeat"):
            n = Node("group" if t == "begin group" else "repeat", row=r, name=g(r, s_idx, "name"),
                     label=proc(g(r, s_idx, lang["label"]), lang))
            stack[-1].children.append(n); stack.append(n)
        elif t in ("end group", "end repeat"):
            if len(stack) > 1: stack.pop()
        elif t == "note":
            stack[-1].children.append(Node("note", row=r, name=g(r, s_idx, "name")))
        elif displayable(r, s_idx, label_cols):
            stack[-1].children.append(Node("q", row=r, name=g(r, s_idx, "name")))
    return root

def count_q(node):
    return sum(1 if ch.kind == "q" else count_q(ch) if ch.kind in ("group", "repeat") else 0 for ch in node.children)

def collect_q(node, out):
    for ch in node.children:
        if ch.kind == "q": out.append(ch)
        elif ch.kind in ("group", "repeat"): collect_q(ch, out)

def is_other_specify(row, s_idx, label_cols):
    if (g(row, s_idx, "type") or "").strip() != "text": return False
    if OTHER_NAME_RE.search(g(row, s_idx, "name") or ""): return True
    lbl = next((g(row, s_idx, c) for c in label_cols if g(row, s_idx, c)), "") or ""
    return bool(OTHER_LBL_RE.search(str(lbl)))

def collect_q_folded(node, s_idx, label_cols):
    raw = []; collect_q(node, raw); out = []
    for q in raw:
        if is_other_specify(q.row, s_idx, label_cols) and out: out[-1] = (out[-1][0], q)
        else: out.append((q, None))
    return out


# ----------------------------------------------------------------------------- emitters
def emit_qtable(doc, qnodes, s_idx, choices, c_idx, lang):
    if not qnodes: return
    tbl = doc.add_table(rows=0, cols=3); tbl.style = "Table Grid"
    for q in qnodes:
        code, text = split_code(proc(g(q.row, s_idx, lang["label"]), lang))
        cells = tbl.add_row().cells
        fill_code_cell(cells[0], code, q.name); fill_cell(cells[1], text)
        fill_cell(cells[2], response_cell(q.row, s_idx, choices, c_idx, lang))
    rest = LANDSCAPE_USABLE_IN - 0.9
    set_col_width(tbl, 0, CODE_COL_W); set_col_width(tbl, 1, Inches(rest * 0.58)); set_col_width(tbl, 2, Inches(rest * 0.42))

def emit_grid(doc, node, s_idx, choices, c_idx, lang):
    cols = []; collect_q(node, cols)
    if not cols: return
    tbl = doc.add_table(rows=0, cols=len(cols)); tbl.style = "Table Grid"
    codes = []; cr = tbl.add_row().cells
    for j, q in enumerate(cols):
        code, _ = split_code(proc(g(q.row, s_idx, lang["label"]), lang))
        codes.append(fill_code_cell(cr[j], code, q.name))
    qr = tbl.add_row().cells
    for j, q in enumerate(cols):
        _, text = split_code(proc(g(q.row, s_idx, lang["label"]), lang))
        resp = response_cell(q.row, s_idx, choices, c_idx, lang)
        fill_cell(qr[j], text + ("\n" + resp if resp else ""))
    for _ in range(3):
        rr = tbl.add_row().cells
        for j in range(len(cols)): fill_cell(rr[j], lang["automatic"] if j == 0 else "")
    flag = {j for j, c in enumerate(codes) if c == CODE_FLAG}
    other_w = Inches(max(0.6, (LANDSCAPE_USABLE_IN - 0.9 * len(flag)) / max(1, len(cols) - len(flag))))
    for j in range(len(cols)): set_col_width(tbl, j, CODE_COL_W if j in flag else other_w)

def parse_if_labels(calc, n=None):
    labs = [clean_html(l).strip() for l in re.findall(r"'([^']*)'", calc or "") if l.strip()]
    return labs[:n] if n else labs

def _qcell_text(q, s_idx, choices, c_idx, lang, other=None):
    code, text = split_code(proc(g(q.row, s_idx, lang["label"]), lang))
    resp = response_cell(q.row, s_idx, choices, c_idx, lang)
    txt = (code + "  " if code else "") + text + ("\n" + resp if resp else "")
    if other is not None:
        ol = proc(g(other.row, s_idx, lang["label"]), lang)
        if ol: txt += "\n" + ol
    return txt

def emit_matrix(doc, node, s_idx, choices, c_idx, lang, cfg, label_cols):
    qs = collect_q_folded(node, s_idx, label_cols)
    if not qs: return
    labels = cfg.get("_labels") or [f"{cfg.get('header', lang['item'])} {i + 1}" for i in range(cfg.get("blank_cols", 5))]
    if cfg.get("orient", "instances_cols") == "instances_rows":
        tbl = doc.add_table(rows=0, cols=1 + len(qs)); tbl.style = "Table Grid"
        hr = tbl.add_row().cells; fill_cell(hr[0], cfg.get("corner", ""), bold=True)
        for j, (q, oth) in enumerate(qs): fill_cell(hr[j + 1], _qcell_text(q, s_idx, choices, c_idx, lang, oth), bold=True)
        for lab in labels: fill_cell(tbl.add_row().cells[0], lab, bold=True)
        set_col_width(tbl, 0, Inches(2.2))
        w = Inches(max(0.7, (LANDSCAPE_USABLE_IN - 2.2) / len(qs)))
        for j in range(1, 1 + len(qs)): set_col_width(tbl, j, w)
    else:
        tbl = doc.add_table(rows=0, cols=1 + len(labels)); tbl.style = "Table Grid"
        hr = tbl.add_row().cells; fill_cell(hr[0], cfg.get("corner", ""), bold=True)
        for j, lab in enumerate(labels): fill_cell(hr[j + 1], lab, bold=True)
        for q, oth in qs: fill_cell(tbl.add_row().cells[0], _qcell_text(q, s_idx, choices, c_idx, lang, oth))
        set_col_width(tbl, 0, Inches(4.0))
        w = Inches(max(0.5, (LANDSCAPE_USABLE_IN - 4.0) / len(labels)))
        for j in range(1, 1 + len(labels)): set_col_width(tbl, j, w)

def render(parent, doc, s_idx, choices, c_idx, lang, buf, label_cols):
    def flush():
        if buf: emit_qtable(doc, buf[:], s_idx, choices, c_idx, lang); buf.clear()
    for node in parent.children:
        if node.kind == "q": buf.append(node)
        elif node.kind == "note":
            flush()
            mk = lang["_note_markers"].get(node.name)
            if mk: add_para(doc, mk, size=10, bold=True)
            add_para(doc, proc(g(node.row, s_idx, lang["label"]), lang))
        elif node.kind == "group":
            if node.label:
                flush(); add_para(doc, node.label, size=13 if SECTION_RE.match(node.label) else 11, bold=True)
            render(node, doc, s_idx, choices, c_idx, lang, buf, label_cols)
        elif node.kind == "repeat":
            mcfg = lang["_matrices"].get(node.name)
            if mcfg:
                flush()
                if node.label: add_para(doc, node.label, size=11, bold=True)
                emit_matrix(doc, node, s_idx, choices, c_idx, lang, mcfg, label_cols)
            elif count_q(node) >= 2:
                flush()
                if node.label: add_para(doc, node.label, size=11, bold=True)
                emit_grid(doc, node, s_idx, choices, c_idx, lang)
            else:
                if node.label: flush(); add_para(doc, node.label, size=11, bold=True)
                render(node, doc, s_idx, choices, c_idx, lang, buf, label_cols)


# ----------------------------------------------------------------------------- document
def _set_landscape(doc):
    sec = doc.sections[0]
    if sec.orientation != WD_ORIENT.LANDSCAPE:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width

def _autofit_window(doc):
    """Set every table to 100% page width with autofit layout; fixed widths act as proportions."""
    for table in doc.tables:
        table.allow_autofit = True; table.autofit = True
        tblPr = table._tbl.tblPr
        tblW = tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW"); style = tblPr.find(qn("w:tblStyle"))
            style.addnext(tblW) if style is not None else tblPr.insert(0, tblW)
        tblW.set(qn("w:type"), "pct"); tblW.set(qn("w:w"), "5000")

def parse_langs(spec):
    """'EN=label:ENG,PT=label' -> [(tag, label_col, hint_col)]"""
    out = []
    for pair in spec.split(","):
        tag, _, col = pair.strip().partition("=")
        if not tag or not col: raise SystemExit(f"bad --langs entry {pair!r}; expected TAG=column")
        hint = col.replace("label", "hint", 1) if col.startswith("label") else "hint"
        out.append((tag.strip(), col.strip(), hint))
    return out

def build_one(form, tag, label_col, hint_col, label_cols, out_path, cfg, autofit=True):
    strings = dict(STRINGS.get(tag) or {})
    strings.update((cfg.get("strings") or {}).get(tag, {}))
    if not strings:
        print(f"WARN  no UI strings for language {tag}; using EN (add config strings.{tag})")
        strings = dict(STRINGS["EN"])
    lang = dict(strings, label=label_col, hint=hint_col, clabel=label_col)
    lang["_subs"] = {v: d[tag] for v, d in (cfg.get("var_labels") or {}).items() if tag in d}
    lang["_note_markers"] = {n: d[tag] for n, d in (cfg.get("note_markers") or {}).items() if tag in d}
    s_idx, s_rows, c_idx, c_rows, title = load(form)
    for c in (label_col, hint_col):
        if c not in s_idx and c == label_col: raise SystemExit(f"column {c!r} not in survey sheet of {form}")
    choices = build_choices(c_idx, c_rows)
    calc_map = {g(r, s_idx, "name"): (g(r, s_idx, "calculation") or "") for r in s_rows if g(r, s_idx, "name")}
    resolved = {}
    for name, m in (cfg.get("matrices") or {}).items():
        c = dict(m); lc = m.get("label_calc")
        if lc and tag in lc: c["_labels"] = parse_if_labels(calc_map.get(lc[tag], ""), m.get("count"))
        if isinstance(m.get("header"), dict): c["header"] = m["header"].get(tag, lang["item"])
        resolved[name] = c
    lang["_matrices"] = resolved
    doc = Document(); _set_landscape(doc)
    banner = (cfg.get("titles") or {}).get(tag) or (str(title) if title else None)
    if banner: add_para(doc, banner, size=13, bold=True, center=True)
    buf = []
    render(parse_tree(s_rows, s_idx, lang, label_cols), doc, s_idx, choices, c_idx, lang, buf, label_cols)
    if buf: emit_qtable(doc, buf[:], s_idx, choices, c_idx, lang)
    if autofit: _autofit_window(doc)
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.save(out_path); print("wrote", out_path)
    return out_path

def export_pdf(docx_path):
    """Export a docx to PDF beside it, using Microsoft Word through COM (Windows with Word installed).
    Returns the PDF path, or None with a printed reason. Elsewhere, LibreOffice does the same job:
    soffice --headless --convert-to pdf <file.docx>"""
    try:
        import win32com.client  # pywin32
    except ImportError:
        print("WARN  --pdf: needs pywin32 + Microsoft Word on Windows; docx written, PDF skipped"
              " (alternative: soffice --headless --convert-to pdf)")
        return None
    pdf_path = os.path.splitext(docx_path)[0] + ".pdf"
    word = None
    try:
        word = win32com.client.DispatchEx("Word.Application"); word.Visible = False
        doc = word.Documents.Open(os.path.abspath(docx_path), False, True)
        doc.SaveAs2(os.path.abspath(pdf_path), 17)   # 17 = wdFormatPDF
        doc.Close(False)
        print("wrote", pdf_path); return pdf_path
    except Exception as e:  # pragma: no cover
        print(f"WARN  --pdf: Word export failed ({e}); docx written, PDF skipped"); return None
    finally:
        if word is not None:
            try: word.Quit()
            except Exception: pass

def generate(form, out_dir, base, langs, cfg=None, suffix="FROMCAPI_v1", autofit=True, pdf=False):
    cfg = cfg or {}
    label_cols = [lc for _, lc, _ in langs]
    outs = []
    for tag, lc, hc in langs:
        p = build_one(form, tag, lc, hc, label_cols, os.path.join(out_dir, tag, f"{base}_{tag}_{suffix}.docx"), cfg, autofit)
        outs.append(p)
        if pdf: export_pdf(p)
    return outs

def load_cfg(path):
    if not path: return {}
    with open(path, encoding="utf-8") as f: return json.load(f)

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("form", nargs="?"); ap.add_argument("--out"); ap.add_argument("--base")
    ap.add_argument("--langs", default="EN=label:ENG,PT=label"); ap.add_argument("--config")
    ap.add_argument("--suffix", default="FROMCAPI_v1"); ap.add_argument("--no-autofit", action="store_true")
    ap.add_argument("--pdf", action="store_true", help="also export each docx to PDF via Microsoft Word (Windows + Word required)")
    ap.add_argument("--batch")
    a = ap.parse_args()
    if a.batch:
        with open(a.batch, encoding="utf-8") as f: b = json.load(f)
        langs = parse_langs(b.get("langs", a.langs)); suffix = b.get("suffix", a.suffix); out = b.get("out", a.out)
        pdf = bool(b.get("pdf", a.pdf))
        if not out: raise SystemExit("batch needs 'out'")
        root = os.path.dirname(os.path.abspath(a.batch))
        if not os.path.isabs(out): out = os.path.normpath(os.path.join(root, out))   # relative to the batch file
        for item in b["forms"]:
            form = item["form"] if os.path.isabs(item["form"]) else os.path.join(root, item["form"])
            cfgp = item.get("config"); cfgp = cfgp if not cfgp or os.path.isabs(cfgp) else os.path.join(root, cfgp)
            if not os.path.exists(form): print(f"MISSING  {form}"); continue
            generate(form, out, item["base"], langs, load_cfg(cfgp), suffix, not a.no_autofit, pdf)
        return
    if not (a.form and a.out and a.base): raise SystemExit("need <form.xlsx> --out <dir> --base <stem>, or --batch")
    generate(a.form, a.out, a.base, parse_langs(a.langs), load_cfg(a.config), a.suffix, not a.no_autofit, a.pdf)

if __name__ == "__main__":
    main()
