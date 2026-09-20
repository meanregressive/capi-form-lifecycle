#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
xlsform_io.py — read an XLSForm as row lists, edit the rows, write it back WITHOUT losing the
workbook's formatting. The helper behind rule G8 ("rebuild sheets from row lists, carry formatting").

Why
---
Editing cells in place with openpyxl is unsafe for XLSForms: `cell(value=None)` does not clear a
cell, and insert/delete rows do not move conditional-formatting ranges, data validations or merged
cells with them. Rebuilding every sheet from a row list is the only way to guarantee no stale cell
reaches the tablet, but a plain rebuild throws away fills, fonts, widths, frozen panes and rules that
teams use to mark translation status or changes. This module does both: it snapshots the formatting
BEFORE the rebuild, keyed by field name (survey), list+value (choices) or column (settings), and
reapplies it AFTER, at the rows' new positions. Rows the edit added inherit the style of the row
above them. Rules and validations are re-anchored by the field names at their ends, and relative row
references inside their formulas are shifted accordingly.

What is carried
---------------
per data row (by key)   font, fill, border, alignment, number format, protection, row height
per column (by header)  width, hidden, header-cell style
per sheet               freeze panes, auto-filter, tab color, zoom, conditional-formatting rules,
                        data validations
rich text in cells      preserved when openpyxl >= 3.1 (loaded with rich_text=True)
NOT carried             merged cells (reported), images, charts, macros, formulas in `version`
                        (the skill requires a literal stamp anyway)

Python API
----------
    from xlsform_io import XLSFormBook
    book = XLSFormBook("base.xlsx")            # snapshot taken here
    rows, hdr = book.rows("survey"), book.header("survey")
    ...edit rows in place (insert / delete / change values)...
    book.save("out.xlsx")                      # rebuild + reapply formatting; returns a report dict

CLI
---
    python xlsform_io.py carry   <base.xlsx> <rebuilt.xlsx> --out <out.xlsx>
        Put base.xlsx's formatting onto a plain rebuilt workbook (values from rebuilt.xlsx).
    python xlsform_io.py inspect <form.xlsx>
        Print a formatting summary (used to check that nothing was lost).
"""
import argparse, io, os, re, sys
from copy import copy
import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.datavalidation import DataValidation

SHEETS = ("survey", "choices", "settings")
_A1 = re.compile(r"(\$?)([A-Za-z]{1,3})(\$?)(\d+)")


# ----------------------------------------------------------------------------- keys
def _s(v):
    return "" if v is None else str(v).strip()

def row_keys(sheet, header, rows):
    """Stable key per data row: survey -> name; choices -> list::value; settings -> 'settings'.
    Repeats (e.g. begin/end group sharing a name) get an occurrence suffix."""
    hl = [h.lower() for h in header]
    def col(*names):
        for n in names:
            if n in hl: return hl.index(n)
        return None
    seen, keys = {}, []
    if sheet == "survey":
        ci = col("name"); ct = col("type")
        for i, r in enumerate(rows):
            base = _s(r[ci]) if ci is not None and ci < len(r) else ""
            if not base: base = f"_{_s(r[ct]) if ct is not None and ct < len(r) else 'row'}_{i}"
            n = seen.get(base, 0); seen[base] = n + 1
            keys.append(base if n == 0 else f"{base}#{n}")
    elif sheet == "choices":
        cl = col("list_name", "list name"); cv = col("value", "name")
        for i, r in enumerate(rows):
            base = f"{_s(r[cl]) if cl is not None and cl < len(r) else ''}::{_s(r[cv]) if cv is not None and cv < len(r) else ''}"
            if base == "::": base = f"_row_{i}"
            n = seen.get(base, 0); seen[base] = n + 1
            keys.append(base if n == 0 else f"{base}#{n}")
    else:
        keys = [f"settings{i}" for i in range(len(rows))]
    return keys


# ----------------------------------------------------------------------------- snapshot
def _cell_style(c):
    return dict(font=copy(c.font), fill=copy(c.fill), border=copy(c.border), alignment=copy(c.alignment),
                number_format=c.number_format, protection=copy(c.protection))

def _apply_style(c, st):
    c.font = st["font"]; c.fill = st["fill"]; c.border = st["border"]; c.alignment = st["alignment"]
    c.number_format = st["number_format"]; c.protection = st["protection"]

def _shift_formula(f, delta):
    if not f or delta == 0 or not isinstance(f, str): return f
    def rep(m):
        if m.group(3) == "$": return m.group(0)                       # absolute row: keep
        return f"{m.group(1)}{m.group(2)}{m.group(3)}{int(m.group(4)) + delta}"
    return _A1.sub(rep, f)

def snapshot_sheet(ws, sheet, header, rows, keys):
    snap = {"row_style": {}, "row_height": {}, "col": {}, "header_style": {},
            "freeze": ws.freeze_panes, "autofilter": ws.auto_filter.ref if ws.auto_filter else None,
            "tab": ws.sheet_properties.tabColor, "zoom": ws.sheet_view.zoomScale,
            "cf": [], "dv": [], "merged": [str(m) for m in ws.merged_cells.ranges]}
    n_hdr = len(header)
    for j, h in enumerate(header, start=1):
        c = ws.cell(row=1, column=j); snap["header_style"][h] = _cell_style(c)
        letter = get_column_letter(j); dim = ws.column_dimensions.get(letter)
        if dim is not None and (dim.width or dim.hidden):
            snap["col"][h] = dict(width=dim.width, hidden=dim.hidden)
    for i, key in enumerate(keys, start=2):
        snap["row_style"][key] = {header[j - 1]: _cell_style(ws.cell(row=i, column=j)) for j in range(1, n_hdr + 1)}
        rd = ws.row_dimensions.get(i)
        if rd is not None and rd.height: snap["row_height"][key] = rd.height
    last = len(rows) + 1
    def anchor(r):                                    # row number -> ("hdr"|"key"|"end", key, overshoot)
        if r <= 1: return ("hdr", None, 0)
        if r >= last: return ("end", None, r - last)
        return ("key", keys[r - 2], 0)
    def hname(c):                                     # column number -> header name (None beyond the header)
        return header[c - 1] if 1 <= c <= n_hdr and header[c - 1] else None
    def span(rng):
        c1, r1, c2, r2 = range_boundaries(str(rng))
        return dict(c1=c1, c2=c2, h1=hname(c1), h2=hname(c2), start=anchor(r1), end=anchor(r2), r1=r1)
    # one item per rule set, keeping ALL its ranges together (a rule on "B1:D1000 F1:F1000" stays one rule)
    for cf in ws.conditional_formatting:
        snap["cf"].append(dict(ranges=[span(r) for r in cf.sqref.ranges], rules=[copy(rule) for rule in cf.rules]))
    for dv in ws.data_validations.dataValidation:
        snap["dv"].append(dict(ranges=[span(r) for r in dv.sqref.ranges], dv=copy(dv)))
    return snap


def reapply_sheet(ws, sheet, header, rows, keys, snap, report):
    n_hdr = len(header); last = len(rows) + 1
    prev_hs = None
    for j, h in enumerate(header, start=1):
        hs = snap["header_style"].get(h) or prev_hs            # a new column's header looks like its neighbor
        if hs: _apply_style(ws.cell(row=1, column=j), hs)
        prev_hs = hs
        if h in snap["col"]:
            dim = ws.column_dimensions[get_column_letter(j)]
            if snap["col"][h]["width"]: dim.width = snap["col"][h]["width"]
            dim.hidden = bool(snap["col"][h]["hidden"])
    inherited = 0; prev_style = None
    for i, key in enumerate(keys, start=2):
        st = snap["row_style"].get(key)
        if st is None:                                  # new row: inherit the row above
            st = prev_style; inherited += 1 if st else 0
        if st:
            for j, h in enumerate(header, start=1):
                if h in st: _apply_style(ws.cell(row=i, column=j), st[h])
        if key in snap["row_height"]: ws.row_dimensions[i].height = snap["row_height"][key]
        prev_style = st
    if snap["freeze"]: ws.freeze_panes = snap["freeze"]
    if snap["autofilter"]:
        c1, r1, c2, r2 = range_boundaries(snap["autofilter"])
        ws.auto_filter.ref = f"{get_column_letter(c1)}{r1}:{get_column_letter(min(c2, n_hdr))}{last}"
    if snap["tab"] is not None: ws.sheet_properties.tabColor = snap["tab"]
    if snap["zoom"]: ws.sheet_view.zoomScale = snap["zoom"]
    pos = {k: i for i, k in enumerate(keys, start=2)}
    hpos = {h: j for j, h in enumerate(header, start=1) if h}
    def resolve(a, fallback):
        kind, key, over = a
        if kind == "hdr": return 1
        if kind == "end": return last + over
        return pos.get(key, fallback)
    def cols(item):                                   # re-anchor columns by header name (columns may have been inserted)
        a = hpos.get(item["h1"]) if item["h1"] else None
        b = hpos.get(item["h2"]) if item["h2"] else None
        if a is None and b is None: return item["c1"], item["c2"]
        if a is None: a = item["c1"]
        if b is None: b = item["c2"] if item["c2"] > n_hdr else b
        if b is None: b = a
        return min(a, b), max(a, b)
    def new_ranges(item):
        """Resolve every range of a rule set at the new positions; return (sqref string, row delta of the first range)."""
        out, delta = [], 0
        for k, sp in enumerate(item["ranges"]):
            r1 = resolve(sp["start"], 2); r2 = max(r1, resolve(sp["end"], last))
            c1, c2 = cols(sp)
            out.append(f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}")
            if k == 0: delta = r1 - sp["r1"]              # formulas are relative to the first range's top-left
        return " ".join(out), delta
    for item in snap["cf"]:
        rng, delta = new_ranges(item)
        for rule in item["rules"]:
            rule = copy(rule)
            if rule.formula: rule.formula = [_shift_formula(f, delta) for f in rule.formula]
            ws.conditional_formatting.add(rng, rule)
        report["cf_rules"] += len(item["rules"])
    for item in snap["dv"]:
        rng, delta = new_ranges(item); old = item["dv"]
        dv = DataValidation(type=old.type, formula1=_shift_formula(old.formula1, delta), formula2=_shift_formula(old.formula2, delta),
                            allow_blank=old.allow_blank, operator=old.operator, showDropDown=old.showDropDown,
                            showErrorMessage=old.showErrorMessage, showInputMessage=old.showInputMessage,
                            error=old.error, errorTitle=old.errorTitle, prompt=old.prompt, promptTitle=old.promptTitle)
        for r in rng.split(" "): dv.add(r)
        ws.add_data_validation(dv); report["dv"] += 1
    if snap["merged"]: report["merged_dropped"] += len(snap["merged"])
    report["rows_inherited"] += inherited


# ----------------------------------------------------------------------------- book
def _load(path):
    try:
        return openpyxl.load_workbook(path, rich_text=True)
    except TypeError:                                   # openpyxl < 3.1
        return openpyxl.load_workbook(path)

def _rows(ws):
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if not rows: return [], []
    header = [_s(h) for h in rows[0]]
    while header and header[-1] == "": header.pop()
    body = [(r + [None] * len(header))[:len(header)] for r in rows[1:]]
    body = [r for r in body if any(v not in (None, "") for v in r)]
    return header, body

class XLSFormBook:
    """Only the XLSForm sheets (survey, choices, settings) are rebuilt; any other sheet (help sheets,
    notes, a template's documentation) is left exactly as loaded."""
    def __init__(self, path):
        self.path = path
        self.wb = _load(path)
        self._header, self._rows, self._snap, self._order, self._index = {}, {}, {}, [], {}
        for idx, name in enumerate(self.wb.sheetnames):
            if name.lower() not in SHEETS: continue
            ws = self.wb[name]; hdr, body = _rows(ws)
            self._header[name] = hdr; self._rows[name] = body; self._order.append(name); self._index[name] = idx
            keys = row_keys(name.lower(), hdr, body)
            self._snap[name] = snapshot_sheet(ws, name.lower(), hdr, body, keys)

    def header(self, sheet): return self._header[sheet]
    def rows(self, sheet): return self._rows[sheet]
    def set_rows(self, sheet, header, rows):
        self._header[sheet] = list(header); self._rows[sheet] = [list(r) for r in rows]
        if sheet not in self._order: self._order.append(sheet); self._index[sheet] = len(self.wb.sheetnames)

    def save(self, out_path, carry_formatting=True):
        """Rebuild the XLSForm sheets from their row lists (remove + create + append), then reapply formatting."""
        report = {"sheets": 0, "cf_rules": 0, "dv": 0, "rows_inherited": 0, "merged_dropped": 0}
        for name in self._order:
            hdr, body = self._header[name], self._rows[name]
            idx = self._index.get(name, len(self.wb.sheetnames))
            if name in self.wb.sheetnames: self.wb.remove(self.wb[name])
            ws = self.wb.create_sheet(name, min(idx, len(self.wb.sheetnames)))
            ws.append(hdr)
            for r in body: ws.append(list(r) + [None] * (len(hdr) - len(r)))
            if carry_formatting and name in self._snap:
                reapply_sheet(ws, name.lower(), hdr, body, row_keys(name.lower(), hdr, body), self._snap[name], report)
            report["sheets"] += 1
        os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
        self.wb.save(out_path)
        return report


# ----------------------------------------------------------------------------- inspect / cli
def summary(path):
    """Formatting counts per sheet, for before/after comparison."""
    wb = _load(path); out = {}
    for name in wb.sheetnames:
        ws = wb[name]
        fills = sum(1 for row in ws.iter_rows() for c in row if c.fill is not None and c.fill.fill_type not in (None, "none"))
        bold = sum(1 for row in ws.iter_rows(min_row=2) for c in row if c.font is not None and c.font.bold)
        colored = sum(1 for row in ws.iter_rows(min_row=2) for c in row if c.font is not None and c.font.color is not None
                      and c.font.color.rgb not in (None, "FF000000") and isinstance(c.font.color.rgb, str))
        widths = sum(1 for d in ws.column_dimensions.values() if d.width)
        hidden = sum(1 for d in ws.column_dimensions.values() if d.hidden)
        cf = sum(len(x.rules) for x in ws.conditional_formatting)
        out[name] = dict(fills=fills, bold=bold, colored_font=colored, widths=widths, hidden_cols=hidden, cf_rules=cf,
                         dv=len(ws.data_validations.dataValidation), freeze=ws.freeze_panes,
                         tab=(ws.sheet_properties.tabColor.rgb if ws.sheet_properties.tabColor is not None else None),
                         merged=len(ws.merged_cells.ranges))
    return out

def carry(base, rebuilt, out):
    book = XLSFormBook(base)
    plain = _load(rebuilt)
    for name in plain.sheetnames:
        if name.lower() in SHEETS:
            hdr, body = _rows(plain[name]); book.set_rows(name, hdr, body)
    for name in list(book._order):
        if name not in plain.sheetnames: book._order.remove(name); book.wb.remove(book.wb[name])
    return book.save(out)

def main():
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")   # only as a script: importers keep their stdout
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    c = sub.add_parser("carry"); c.add_argument("base"); c.add_argument("rebuilt"); c.add_argument("--out", required=True)
    i = sub.add_parser("inspect"); i.add_argument("form")
    a = ap.parse_args()
    if a.cmd == "carry":
        rep = carry(a.base, a.rebuilt, a.out); print(f"wrote {a.out}  {rep}")
    else:
        for name, d in summary(a.form).items(): print(f"{name}: {d}")

if __name__ == "__main__":
    main()
